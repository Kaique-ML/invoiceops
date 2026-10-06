from __future__ import annotations

import asyncio
import csv
import hashlib
import hmac
import io
import json
import logging
import re
import secrets
import shutil
import unicodedata
from collections.abc import AsyncIterator, Awaitable, Callable, Generator
from contextlib import asynccontextmanager
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from time import perf_counter
from typing import Annotated
from uuid import UUID, uuid4

import redis
from argon2 import PasswordHasher
from argon2.exceptions import VerifyMismatchError
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from sqlalchemy import func, select, text
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import IntegrityError, OperationalError, SQLAlchemyError
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from invoiceops.domain.exports import export_cell
from invoiceops.domain.schemas import ApprovedRecordPayload, InvoiceExtraction, PageText
from invoiceops.domain.validation import validate_extraction
from invoiceops.extraction.documents import detect_document_type
from invoiceops.persistence.database import make_session_factory
from invoiceops.persistence.models import (
    AuditEvent,
    Delivery,
    Document,
    IntegrationReceipt,
    OutboxEvent,
    ProcessingAttempt,
    Revision,
    User,
    utc_now,
)
from invoiceops.services.processing import (
    allowed_demo_hashes,
    normalize_extraction,
)
from invoiceops.services.queue import recover_stale_processing
from invoiceops.services.storage import LocalDocumentStorage
from invoiceops.services.webhook import ConfiguredWebhookSink, WebhookConfigurationError
from invoiceops.settings import Settings

logger = logging.getLogger("invoiceops.api")
BLOCKING_ISSUE_CODES = {
    "unverified_evidence",
    "missing_evidence",
    "unsupported_currency",
    "ambiguous_amount",
    "ambiguous_date",
    "total_mismatch",
}
settings = Settings.from_environment()
settings.validate()
session_factory = make_session_factory(settings)
storage = LocalDocumentStorage(settings.upload_dir, max_bytes=settings.max_upload_bytes)
templates = Jinja2Templates(directory="web/templates")
password_hasher = PasswordHasher()


def _current_session() -> Session:
    return session_factory()


def get_db() -> Generator[Session, None, None]:
    with _current_session() as session:
        yield session


def current_user(
    request: Request,
    session: Annotated[Session, Depends(get_db)],
) -> User:
    raw_id = request.session.get("user_id")
    try:
        user_id = UUID(raw_id) if isinstance(raw_id, str) else None
    except ValueError:
        user_id = None
    user = session.get(User, user_id) if user_id else None
    if user is None or not user.active:
        raise HTTPException(status_code=401, detail="Authentication required.")
    return user


def require_reviewer(user: Annotated[User, Depends(current_user)]) -> User:
    if user.role not in {"admin", "reviewer"}:
        raise HTTPException(status_code=403, detail="Reviewer or administrator role required.")
    return user


def require_admin(user: Annotated[User, Depends(current_user)]) -> User:
    if user.role != "admin":
        raise HTTPException(status_code=403, detail="Administrator role required.")
    return user


def _csrf(request: Request, supplied: str | None) -> None:
    expected = request.session.get("csrf_token")
    if (
        not isinstance(expected, str)
        or not supplied
        or not secrets.compare_digest(expected, supplied)
    ):
        raise HTTPException(status_code=403, detail="CSRF token missing or invalid.")


def _csrf_token(request: Request) -> str:
    value = request.session.get("csrf_token")
    if not isinstance(value, str):
        value = secrets.token_urlsafe(32)
        request.session["csrf_token"] = value
    return value


def _safe_filename(filename: str | None) -> str:
    normalized = (filename or "upload").replace("\\", "/")
    base = PurePosixPath(normalized).name
    clean = "".join(char for char in base if unicodedata.category(char) not in {"Cc", "Cf"})
    return clean[:255] or "upload"


def _latest_revision(session: Session, document_id: UUID) -> Revision | None:
    return session.scalar(
        select(Revision)
        .where(Revision.document_id == document_id)
        .order_by(Revision.number.desc())
        .limit(1)
    )


def _get_document(
    session: Session,
    user: User,
    document_id: UUID,
    *,
    lock: bool = False,
) -> Document:
    query = select(Document).where(
        Document.id == document_id,
        Document.workspace_id == user.workspace_id,
    )
    if lock:
        query = query.with_for_update()
    document = session.scalar(query)
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found.")
    return document


def _create_app() -> FastAPI:
    @asynccontextmanager
    async def lifespan(_: FastAPI) -> AsyncIterator[None]:
        task = asyncio.create_task(_outbox_loop())
        yield
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass

    application = FastAPI(
        title="InvoiceOps",
        version="0.1.0",
        description="Document intake and human review demonstration.",
        lifespan=lifespan,
    )
    application.add_middleware(
        SessionMiddleware,
        secret_key=settings.session_secret,
        session_cookie="invoiceops_session",
        max_age=8 * 60 * 60,
        same_site="lax",
        https_only=settings.app_env == "production",
    )
    application.mount("/static", StaticFiles(directory="web/static"), name="static")

    @application.middleware("http")
    async def request_context(
        request: Request, call_next: Callable[[Request], Awaitable[Response]]
    ) -> Response:
        started = perf_counter()
        supplied_id = request.headers.get("X-Request-ID", "")
        correlation_id = (
            supplied_id
            if 1 <= len(supplied_id) <= 64
            and all(char.isalnum() or char in "._-" for char in supplied_id)
            else str(uuid4())
        )
        request.state.correlation_id = correlation_id
        try:
            response = await call_next(request)
        except Exception:
            logger.exception(
                json.dumps(
                    {
                        "event": "request_failed",
                        "request_id": correlation_id,
                        "method": request.method,
                        "path": request.url.path,
                        "duration_ms": round((perf_counter() - started) * 1000, 2),
                    }
                )
            )
            raise
        response.headers["X-Request-ID"] = correlation_id
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "same-origin"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; img-src 'self' data:; style-src 'self'; "
            "form-action 'self'; frame-ancestors 'none'; object-src 'none'; base-uri 'self'"
        )
        logger.info(
            json.dumps(
                {
                    "event": "request_completed",
                    "request_id": correlation_id,
                    "method": request.method,
                    "path": request.url.path,
                    "status_code": response.status_code,
                    "duration_ms": round((perf_counter() - started) * 1000, 2),
                }
            )
        )
        return response

    @application.get("/", include_in_schema=False)
    def home() -> RedirectResponse:
        return RedirectResponse("/documents", status_code=303)

    @application.get("/health/live")
    def live() -> dict[str, str]:
        return {"status": "alive"}

    @application.get("/health/ready")
    def ready(session: Annotated[Session, Depends(get_db)]) -> Response:
        status: dict[str, bool | None] = {
            "database": False,
            "redis": False,
            "ocr": False,
            "model_configured": None,
        }
        try:
            session.execute(text("select 1"))
            status["database"] = True
        except OperationalError:
            pass
        try:
            status["redis"] = bool(
                redis.Redis.from_url(settings.redis_url, socket_timeout=1).ping()
            )
        except (OSError, redis.RedisError):
            status["redis"] = False
        status["ocr"] = shutil.which("tesseract") is not None
        if settings.mode == "real":
            status["model_configured"] = bool(settings.ollama_model)
        healthy = bool(status["database"] and status["redis"] and status["ocr"])
        if settings.mode == "real":
            healthy = healthy and bool(status["model_configured"])
        return Response(
            content=json.dumps({"status": "ready" if healthy else "not_ready", **status}),
            media_type="application/json",
            status_code=200 if healthy else 503,
        )

    @application.get("/metrics", include_in_schema=False)
    def metrics(session: Annotated[Session, Depends(get_db)]) -> Response:
        return Response(content=_metrics_text(session), media_type="text/plain; version=0.0.4")

    @application.get("/login", response_class=HTMLResponse, include_in_schema=False)
    def login_page(request: Request) -> HTMLResponse:
        return templates.TemplateResponse(
            request,
            "login.html",
            {"csrf_token": _csrf_token(request), "error": None},
        )

    @application.post("/login", response_class=HTMLResponse, include_in_schema=False)
    async def login(
        request: Request,
        session: Annotated[Session, Depends(get_db)],
    ) -> Response:
        form = await request.form()
        _csrf(request, str(form.get("csrf_token", "")))
        username = str(form.get("username", "")).strip()[:128]
        password = str(form.get("password", ""))
        matches = list(session.scalars(select(User).where(User.username == username, User.active)))
        valid = len(matches) == 1
        user = matches[0] if valid else None
        if user is not None:
            try:
                password_hasher.verify(user.password_hash, password)
            except VerifyMismatchError:
                valid = False
        if not valid or user is None:
            return templates.TemplateResponse(
                request,
                "login.html",
                {"csrf_token": _csrf_token(request), "error": "Usuário ou senha inválidos."},
                status_code=401,
            )
        request.session.clear()
        request.session["user_id"] = str(user.id)
        request.session["csrf_token"] = secrets.token_urlsafe(32)
        return RedirectResponse("/documents", status_code=303)

    @application.post("/logout", include_in_schema=False)
    async def logout(request: Request) -> Response:
        form = await request.form()
        _csrf(request, str(form.get("csrf_token", "")))
        request.session.clear()
        return RedirectResponse("/login", status_code=303)

    @application.get("/documents", response_class=HTMLResponse, include_in_schema=False)
    def document_list(
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        q: str = "",
        state: str = "",
        page: int = 1,
    ) -> Response:
        user = _user_or_none(request, session)
        if user is None:
            return RedirectResponse("/login", status_code=303)
        page = max(1, page)
        query = select(Document).where(Document.workspace_id == user.workspace_id)
        if q.strip():
            query = query.where(Document.original_name.ilike(f"%{q.strip()[:80]}%"))
        if state in {"queued", "processing", "completed", "failed"}:
            query = query.where(Document.processing_status == state)
        records = list(
            session.scalars(
                query.order_by(Document.created_at.desc()).offset((page - 1) * 25).limit(25)
            )
        )
        return templates.TemplateResponse(
            request,
            "documents.html",
            {
                "documents": records,
                "user": user,
                "workspace": user.workspace.name,
                "mode": settings.mode,
                "csrf_token": _csrf_token(request),
                "page": page,
                "query": q,
                "state": state,
            },
        )

    @application.post("/upload", include_in_schema=False)
    async def upload_form(
        request: Request,
        file: Annotated[UploadFile, File()],
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_reviewer)],
        csrf_token: Annotated[str, Form()],
    ) -> Response:
        _csrf(request, csrf_token)
        document, duplicate = await _accept_upload(file, session, user)
        if duplicate:
            request.session["notice"] = "Arquivo idêntico já recebido neste workspace."
        else:
            request.session["notice"] = "Arquivo aceito para processamento assíncrono."
        return RedirectResponse(f"/documents/{document.id}", status_code=303)

    @application.get(
        "/documents/{document_id}", response_class=HTMLResponse, include_in_schema=False
    )
    def document_detail(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
    ) -> Response:
        user = _user_or_none(request, session)
        if user is None:
            return RedirectResponse("/login", status_code=303)
        document = _get_document(session, user, document_id)
        revision = _latest_revision(session, document.id)
        audit = list(
            session.scalars(
                select(AuditEvent)
                .where(AuditEvent.document_id == document.id)
                .order_by(AuditEvent.created_at.desc())
                .limit(100)
            )
        )
        deliveries = list(
            session.scalars(
                select(Delivery)
                .where(
                    Delivery.document_id == document.id,
                    Delivery.workspace_id == user.workspace_id,
                )
                .order_by(Delivery.created_at.desc())
            )
        )
        return templates.TemplateResponse(
            request,
            "detail.html",
            {
                "document": document,
                "revision": revision,
                "audit": audit,
                "deliveries": deliveries,
                "user": user,
                "workspace": user.workspace.name,
                "mode": settings.mode,
                "csrf_token": _csrf_token(request),
                "notice": request.session.pop("notice", None),
                "webhook_configured": bool(settings.webhook_url and settings.webhook_secret),
                "has_blocking_issues": bool(
                    revision
                    and any(issue.get("code") in BLOCKING_ISSUE_CODES for issue in revision.issues)
                ),
            },
        )

    @application.get("/documents/{document_id}/file", include_in_schema=False)
    def original_file(
        document_id: UUID,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
    ) -> FileResponse:
        document = _get_document(session, user, document_id)
        path = (storage.root / document.storage_key).resolve()
        if storage.root not in path.parents or not path.is_file():
            raise HTTPException(status_code=404, detail="Original file is unavailable.")
        return FileResponse(
            path,
            media_type=document.content_type,
            filename=document.original_name,
            content_disposition_type=(
                "inline" if document.content_type in {"image/png", "image/jpeg"} else "attachment"
            ),
            headers={"X-Content-Type-Options": "nosniff"},
        )

    @application.post("/documents/{document_id}/review", include_in_schema=False)
    async def review_form(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_reviewer)],
    ) -> Response:
        form = await request.form()
        _csrf(request, str(form.get("csrf_token", "")))
        expected_version = int(str(form.get("version", "0")))
        changes = {
            field: str(form.get(field, "")).strip() or None
            for field in (
                "supplier",
                "tax_identifier",
                "document_number",
                "issue_date",
                "currency",
                "subtotal",
                "discounts",
                "taxes",
                "shipping",
                "total",
            )
        }
        _save_review(
            session,
            user,
            document_id,
            expected_version,
            changes,
            str(form.get("reason", "")).strip(),
        )
        return RedirectResponse(f"/documents/{document_id}", status_code=303)

    @application.post("/documents/{document_id}/approve", include_in_schema=False)
    async def approve_form(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_reviewer)],
    ) -> Response:
        form = await request.form()
        _csrf(request, str(form.get("csrf_token", "")))
        _approve(session, user, document_id, int(str(form.get("version", "0"))))
        return RedirectResponse(f"/documents/{document_id}", status_code=303)

    @application.post("/documents/{document_id}/reject", include_in_schema=False)
    async def reject_form(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_reviewer)],
    ) -> Response:
        form = await request.form()
        _csrf(request, str(form.get("csrf_token", "")))
        document = _get_document(session, user, document_id, lock=True)
        expected_version = int(str(form.get("version", "0")))
        if document.version != expected_version:
            raise HTTPException(
                status_code=409, detail="Document changed; reload before rejecting."
            )
        if document.processing_status != "completed":
            raise HTTPException(status_code=409, detail="Incomplete processing cannot be rejected.")
        document.review_status = "rejected"
        document.version += 1
        session.add(
            AuditEvent(
                document_id=document.id,
                actor_id=user.id,
                action="rejected",
                reason=str(form.get("reason", "")).strip()[:500] or None,
            )
        )
        session.commit()
        return RedirectResponse(f"/documents/{document_id}", status_code=303)

    @application.post("/documents/{document_id}/reprocess", include_in_schema=False)
    async def reprocess_form(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_reviewer)],
    ) -> Response:
        form = await request.form()
        _csrf(request, str(form.get("csrf_token", "")))
        document = _get_document(session, user, document_id, lock=True)
        _request_reprocess(session, user, document)
        return RedirectResponse(f"/documents/{document_id}", status_code=303)

    @application.post("/documents/{document_id}/deliver", include_in_schema=False)
    async def deliver_form(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_admin)],
    ) -> Response:
        form = await request.form()
        _csrf(request, str(form.get("csrf_token", "")))
        delivery, duplicate = _request_delivery(session, user, document_id)
        request.session["notice"] = (
            "Esta revisão já possui uma entrega registrada."
            if duplicate
            else "Entrega persistida na fila para processamento assíncrono."
        )
        return RedirectResponse(f"/documents/{document_id}", status_code=303)

    @application.get("/api/v1/documents")
    def api_list_documents(
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
        page: int = 1,
        processing_status: str = "",
    ) -> dict[str, object]:
        page = max(1, page)
        query = select(Document).where(Document.workspace_id == user.workspace_id)
        if processing_status in {"queued", "processing", "completed", "failed"}:
            query = query.where(Document.processing_status == processing_status)
        rows = list(
            session.scalars(
                query.order_by(Document.created_at.desc()).offset((page - 1) * 25).limit(25)
            )
        )
        return {"page": page, "items": [_document_json(row) for row in rows]}

    @application.get("/api/v1/documents/{document_id}")
    def api_document(
        document_id: UUID,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
    ) -> dict[str, object]:
        document = _get_document(session, user, document_id)
        revision = _latest_revision(session, document.id)
        return {
            "document": _document_json(document),
            "revision": _revision_json(revision) if revision else None,
        }

    @application.get("/api/v1/documents/{document_id}/revisions")
    def api_revisions(
        document_id: UUID,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
    ) -> list[dict[str, object]]:
        _get_document(session, user, document_id)
        revisions = list(
            session.scalars(
                select(Revision)
                .where(Revision.document_id == document_id)
                .order_by(Revision.number)
            )
        )
        return [_revision_json(revision) for revision in revisions]

    @application.get("/api/v1/documents/{document_id}/deliveries")
    def api_deliveries(
        document_id: UUID,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
    ) -> list[dict[str, object]]:
        _get_document(session, user, document_id)
        rows = list(
            session.scalars(
                select(Delivery)
                .where(
                    Delivery.document_id == document_id,
                    Delivery.workspace_id == user.workspace_id,
                )
                .order_by(Delivery.created_at)
            )
        )
        return [_delivery_json(row) for row in rows]

    @application.get("/api/v1/documents/{document_id}/audit")
    def api_audit(
        document_id: UUID,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
    ) -> list[dict[str, object]]:
        _get_document(session, user, document_id)
        events = list(
            session.scalars(
                select(AuditEvent)
                .where(AuditEvent.document_id == document_id)
                .order_by(AuditEvent.created_at)
            )
        )
        return [
            {
                "id": str(event.id),
                "actor_id": str(event.actor_id),
                "action": event.action,
                "field": event.field_name,
                "old_value": event.old_value,
                "new_value": event.new_value,
                "reason": event.reason,
                "created_at": event.created_at.isoformat(),
            }
            for event in events
        ]

    @application.post("/api/v1/documents/{document_id}/approve")
    def api_approve(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_reviewer)],
        expected_version: int,
        x_csrf_token: Annotated[str | None, Header()] = None,
    ) -> dict[str, str]:
        _csrf(request, x_csrf_token)
        _approve(session, user, document_id, expected_version)
        return {"status": "approved"}

    @application.post("/api/v1/documents/{document_id}/deliver", status_code=202)
    def api_deliver(
        document_id: UUID,
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_admin)],
        x_csrf_token: Annotated[str | None, Header()] = None,
    ) -> dict[str, object]:
        _csrf(request, x_csrf_token)
        delivery, duplicate = _request_delivery(session, user, document_id)
        return {"delivery": _delivery_json(delivery), "duplicate": duplicate}

    @application.post("/api/v1/documents", status_code=202)
    async def api_upload(
        request: Request,
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(require_reviewer)],
        file: Annotated[UploadFile, File()],
        x_csrf_token: Annotated[str | None, Header()] = None,
    ) -> dict[str, object]:
        _csrf(request, x_csrf_token)
        document, duplicate = await _accept_upload(file, session, user)
        return {"document": _document_json(document), "duplicate": duplicate}

    @application.get("/api/v1/exports.csv")
    def export_csv(
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
    ) -> Response:
        rows = _approved_rows(session, user)
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=_export_columns())
        writer.writeheader()
        for row in rows:
            writer.writerow({key: export_cell(value) for key, value in row.items()})
        return Response(
            buffer.getvalue(),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": 'attachment; filename="invoiceops-approved.csv"'},
        )

    @application.get("/api/v1/exports.xlsx")
    def export_xlsx(
        session: Annotated[Session, Depends(get_db)],
        user: Annotated[User, Depends(current_user)],
    ) -> Response:
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Approved records")
        columns = _export_columns()
        sheet.append(columns)
        for row in _approved_rows(session, user):
            sheet.append([export_cell(row[key]) for key in columns])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return Response(
            buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="invoiceops-approved.xlsx"'},
        )

    @application.post("/integration/test/receive")
    def receive_test_webhook(
        payload: ApprovedRecordPayload,
        session: Annotated[Session, Depends(get_db)],
        authorization: Annotated[str | None, Header()] = None,
        idempotency_key: Annotated[str | None, Header(alias="Idempotency-Key")] = None,
    ) -> dict[str, object]:
        if settings.app_env != "development":
            raise HTTPException(
                status_code=404, detail="Test receiver is disabled outside development."
            )
        if not settings.webhook_secret or len(settings.webhook_secret) < 32:
            raise HTTPException(
                status_code=503,
                detail="Configure a local receiver secret of at least 32 characters.",
            )
        supplied = authorization.removeprefix("Bearer ") if authorization else ""
        if not hmac.compare_digest(supplied, settings.webhook_secret):
            raise HTTPException(status_code=401, detail="Webhook authentication failed.")
        if not idempotency_key or not re.fullmatch(r"[A-Za-z0-9._:-]{1,160}", idempotency_key):
            raise HTTPException(status_code=400, detail="A valid Idempotency-Key is required.")
        if idempotency_key != f"invoiceops-v1-{payload.event_id}":
            raise HTTPException(
                status_code=400, detail="Idempotency-Key does not match the event identifier."
            )
        body = payload.model_dump(mode="json")
        canonical = json.dumps(body, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        digest = hashlib.sha256(canonical.encode("utf-8")).hexdigest()
        statement = (
            pg_insert(IntegrationReceipt)
            .values(
                idempotency_key=idempotency_key,
                payload_hash=digest,
                payload=body,
            )
            .on_conflict_do_nothing(index_elements=[IntegrationReceipt.idempotency_key])
            .returning(IntegrationReceipt.id)
        )
        inserted = session.execute(statement).scalar_one_or_none()
        existing = session.scalar(
            select(IntegrationReceipt).where(IntegrationReceipt.idempotency_key == idempotency_key)
        )
        if existing is not None and existing.payload_hash != digest:
            session.rollback()
            raise HTTPException(
                status_code=409,
                detail="The idempotency key was previously used with a different payload.",
            )
        session.commit()
        return {
            "status": "processed",
            "replayed": inserted is None,
            "receipt_id": str(existing.id) if existing is not None else None,
        }

    return application


def _user_or_none(request: Request, session: Session) -> User | None:
    raw_id = request.session.get("user_id")
    try:
        user_id = UUID(raw_id) if isinstance(raw_id, str) else None
    except ValueError:
        user_id = None
    user = session.get(User, user_id) if user_id else None
    return user if user is not None and user.active else None


async def _accept_upload(file: UploadFile, session: Session, user: User) -> tuple[Document, bool]:
    content = await file.read(settings.max_upload_bytes + 1)
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > settings.max_upload_bytes:
        raise HTTPException(
            status_code=413, detail="The uploaded file exceeds the configured byte limit."
        )
    filename = _safe_filename(file.filename)
    try:
        content_type = detect_document_type(content, filename)
    except ValueError as exc:
        raise HTTPException(status_code=415, detail=str(exc)) from exc
    digest = hashlib.sha256(content).hexdigest()
    if settings.mode == "demo" and digest not in allowed_demo_hashes():
        raise HTTPException(
            status_code=400,
            detail=(
                "Demo mode accepts only the bundled synthetic samples. "
                "Configure real mode for other files."
            ),
        )
    existing = session.scalar(
        select(Document).where(
            Document.workspace_id == user.workspace_id,
            Document.sha256 == digest,
        )
    )
    if existing is not None:
        session.add(
            AuditEvent(
                document_id=existing.id,
                actor_id=user.id,
                action="exact_duplicate_detected",
            )
        )
        session.commit()
        return existing, True

    document_id = uuid4()
    try:
        storage_key = storage.store(user.workspace_id, document_id, content)
    except ValueError as exc:
        raise HTTPException(status_code=413, detail=str(exc)) from exc
    document = Document(
        id=document_id,
        workspace_id=user.workspace_id,
        uploaded_by=user.id,
        original_name=filename,
        content_type=content_type,
        sha256=digest,
        size_bytes=len(content),
        storage_key=storage_key,
        processing_status="queued",
        review_status="pending",
    )
    session.add(document)
    session.add(
        OutboxEvent(
            workspace_id=user.workspace_id,
            document_id=document_id,
            event_type="process",
            state="pending",
        )
    )
    try:
        session.commit()
    except IntegrityError as exc:
        session.rollback()
        storage.delete(storage_key)
        duplicate = session.scalar(
            select(Document).where(
                Document.workspace_id == user.workspace_id,
                Document.sha256 == digest,
            )
        )
        if duplicate is not None:
            session.add(
                AuditEvent(
                    document_id=duplicate.id,
                    actor_id=user.id,
                    action="exact_duplicate_detected",
                )
            )
            session.commit()
            return duplicate, True
        raise HTTPException(
            status_code=409, detail="Upload conflicted with another request; retry safely."
        ) from exc
    except SQLAlchemyError as exc:
        session.rollback()
        storage.delete(storage_key)
        raise HTTPException(
            status_code=503, detail="Upload intent could not be committed; retry safely."
        ) from exc
    session.refresh(document)
    return document, False


def _approve(session: Session, user: User, document_id: UUID, expected_version: int) -> None:
    document = _get_document(session, user, document_id, lock=True)
    if document.version != expected_version:
        raise HTTPException(status_code=409, detail="Document changed; reload before approving.")
    if document.processing_status != "completed":
        raise HTTPException(
            status_code=409, detail="Only fully processed documents can be approved."
        )
    revision = _latest_revision(session, document.id)
    if revision is None:
        raise HTTPException(status_code=409, detail="No completed extraction revision exists.")
    if revision.approved_at is not None or document.review_status == "approved":
        raise HTTPException(status_code=409, detail="This revision has already been approved.")
    blocking = [
        issue for issue in revision.issues if issue.get("code") in BLOCKING_ISSUE_CODES
    ]
    if blocking:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "Resolve blocking validation issues before approval.",
                "issues": blocking,
            },
        )
    revision.approved_by = user.id
    revision.approved_at = utc_now()
    document.review_status = "approved"
    document.version += 1
    session.add(
        AuditEvent(
            document_id=document.id,
            revision_id=revision.id,
            actor_id=user.id,
            action="approved",
        )
    )
    session.commit()


def _save_review(
    session: Session,
    user: User,
    document_id: UUID,
    expected_version: int,
    changes: dict[str, str | None],
    reason: str,
) -> None:
    document = _get_document(session, user, document_id, lock=True)
    if document.version != expected_version:
        raise HTTPException(
            status_code=409, detail="Document changed; reload before saving corrections."
        )
    if document.processing_status != "completed" or document.review_status == "rejected":
        raise HTTPException(status_code=409, detail="This document is not available for review.")
    previous = _latest_revision(session, document.id)
    if previous is None:
        raise HTTPException(status_code=409, detail="No extraction revision is available.")
    approval_was_revoked = previous.approved_at is not None
    normalized = dict(previous.normalized)
    now = utc_now().isoformat()
    provenance = dict(normalized.get("field_provenance", {}))
    audit_events: list[AuditEvent] = []
    for field, value in changes.items():
        old_value = normalized.get(field)
        if old_value == value:
            continue
        normalized[field] = value
        provenance[field] = {
            "origin": "human-review",
            "user_id": str(user.id),
            "at": now,
            "reason": reason[:500] or None,
        }
        audit_events.append(
            AuditEvent(
                document_id=document.id,
                revision_id=None,
                actor_id=user.id,
                action="field_corrected",
                field_name=field,
                old_value=str(old_value) if old_value is not None else None,
                new_value=value,
                reason=reason[:500] or None,
            )
        )
    if not audit_events:
        raise HTTPException(status_code=400, detail="No field values changed.")
    if approval_was_revoked:
        previous.approved_at = None
        previous.approved_by = None
        audit_events.append(
            AuditEvent(
                document_id=document.id,
                revision_id=previous.id,
                actor_id=user.id,
                action="approval_revoked",
                reason=(reason[:500] or "Approved data was corrected."),
            )
        )
    normalized["field_provenance"] = provenance
    current_values = dict(previous.extracted)
    for field in (
        "document_type",
        "supplier",
        "tax_identifier",
        "document_number",
        "issue_date",
        "currency",
        "subtotal",
        "discounts",
        "taxes",
        "shipping",
        "total",
        "items",
        "tax_included_in_total",
    ):
        if field in normalized:
            current_values[field] = normalized[field]
    corrected = InvoiceExtraction.model_validate(current_values)
    normalized_values, normalize_issues = normalize_extraction(corrected)
    normalized.update(normalized_values)
    source_pages = [PageText.model_validate(page) for page in previous.source_pages]
    issues = (
        validate_extraction(
            corrected,
            source_pages,
            manually_verified_fields=set(provenance),
        )
        + normalize_issues
    )
    for event in audit_events:
        session.add(event)
    revision = Revision(
        document_id=document.id,
        number=previous.number + 1,
        kind="review",
        source_pages=previous.source_pages,
        extracted=previous.extracted,
        normalized=normalized,
        issues=issues,
        mode=previous.mode,
        model_name=previous.model_name,
        prompt_version=previous.prompt_version,
    )
    session.add(revision)
    session.flush()
    for event in audit_events:
        event.revision_id = revision.id
    document.review_status = "needs_review" if issues else "pending"
    if approval_was_revoked:
        document.delivery_status = "not_requested"
    document.version += 1
    session.commit()


def _request_reprocess(session: Session, user: User, document: Document) -> None:
    if document.processing_status in {"queued", "processing"}:
        raise HTTPException(status_code=409, detail="Document is already queued or processing.")
    document.processing_status = "queued"
    document.review_status = "pending"
    document.processing_started_at = None
    document.claim_token = None
    document.last_error = None
    document.version += 1
    session.add(
        OutboxEvent(
            workspace_id=user.workspace_id,
            document_id=document.id,
            event_type="process",
            state="pending",
        )
    )
    session.add(
        AuditEvent(
            document_id=document.id,
            actor_id=user.id,
            action="reprocess_requested",
        )
    )
    session.commit()


def _request_delivery(
    session: Session,
    user: User,
    document_id: UUID,
) -> tuple[Delivery, bool]:
    document = _get_document(session, user, document_id, lock=True)
    if document.processing_status != "completed" or document.review_status != "approved":
        raise HTTPException(
            status_code=409, detail="Only a completed, human-approved document can be delivered."
        )
    revision = _latest_revision(session, document.id)
    if revision is None or revision.approved_at is None:
        raise HTTPException(status_code=409, detail="No approved revision snapshot is available.")
    try:
        ConfiguredWebhookSink(settings)
    except WebhookConfigurationError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    destination_key = hashlib.sha256(settings.webhook_url.encode("utf-8")).hexdigest()[:32]
    existing = session.scalar(
        select(Delivery).where(
            Delivery.revision_id == revision.id,
            Delivery.destination_key == destination_key,
        )
    )
    if existing is not None:
        return existing, True

    delivery_id = uuid4()
    fields = {
        key: revision.normalized.get(key)
        for key in (
            "document_type",
            "supplier",
            "tax_identifier",
            "document_number",
            "issue_date",
            "currency",
            "subtotal",
            "discounts",
            "taxes",
            "shipping",
            "total",
            "items",
        )
    }
    payload = {
        "schema_version": "invoiceops.approved.v1",
        "event_id": str(delivery_id),
        "document_id": str(document.id),
        "revision_id": str(revision.id),
        "approved_at": revision.approved_at.isoformat(),
        "document_sha256": document.sha256,
        "fields": fields,
    }
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    delivery = Delivery(
        id=delivery_id,
        workspace_id=user.workspace_id,
        document_id=document.id,
        revision_id=revision.id,
        destination_key=destination_key,
        idempotency_key=f"invoiceops-v1-{delivery_id}",
        payload_version="v1",
        payload=payload,
        payload_hash=hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
        state="pending",
    )
    session.add(delivery)
    session.add(
        OutboxEvent(
            workspace_id=user.workspace_id,
            document_id=document.id,
            delivery_id=delivery.id,
            event_type="delivery",
            state="pending",
        )
    )
    document.delivery_status = "pending"
    session.add(
        AuditEvent(
            document_id=document.id,
            revision_id=revision.id,
            actor_id=user.id,
            action="delivery_requested",
        )
    )
    session.commit()
    session.refresh(delivery)
    return delivery, False


def _approved_rows(session: Session, user: User) -> list[dict[str, object]]:
    approved_documents = list(
        session.scalars(
            select(Document)
            .where(
                Document.workspace_id == user.workspace_id,
                Document.review_status == "approved",
            )
            .order_by(Document.created_at)
        )
    )
    records: list[tuple[Document, Revision]] = []
    for document in approved_documents:
        revision = _latest_revision(session, document.id)
        if revision is not None and revision.approved_at is not None:
            records.append((document, revision))
    rows: list[dict[str, object]] = []
    for document, revision in records:
        data = revision.normalized
        rows.append(
            {
                "document_id": str(document.id),
                "revision_id": str(revision.id),
                "document_name": document.original_name,
                "sha256": document.sha256,
                "document_type": data.get("document_type"),
                "supplier": data.get("supplier"),
                "tax_identifier": data.get("tax_identifier"),
                "document_number": data.get("document_number"),
                "issue_date": data.get("issue_date"),
                "currency": data.get("currency"),
                "subtotal": _decimal_export(data.get("subtotal")),
                "discounts": _decimal_export(data.get("discounts")),
                "taxes": _decimal_export(data.get("taxes")),
                "shipping": _decimal_export(data.get("shipping")),
                "total": _decimal_export(data.get("total")),
                "approved_at": revision.approved_at.isoformat() if revision.approved_at else None,
            }
        )
    return rows


def _metrics_text(session: Session) -> str:
    lines = [
        "# HELP invoiceops_documents_total Documents accepted by current state.",
        "# TYPE invoiceops_documents_total gauge",
    ]
    total_documents = session.scalar(select(func.count()).select_from(Document)) or 0
    lines.append(f'invoiceops_documents_total{{processing_status="all"}} {total_documents}')
    for state, count in session.execute(
        select(Document.processing_status, func.count()).group_by(Document.processing_status)
    ):
        lines.append(f'invoiceops_documents_total{{processing_status="{state}"}} {count}')
    lines.extend(
        [
            "# HELP invoiceops_reviews_total Documents by current review state.",
            "# TYPE invoiceops_reviews_total gauge",
        ]
    )
    for state, count in session.execute(
        select(Document.review_status, func.count()).group_by(Document.review_status)
    ):
        lines.append(f'invoiceops_reviews_total{{review_status="{state}"}} {count}')
    lines.extend(
        [
            "# HELP invoiceops_deliveries_total Delivery records by state.",
            "# TYPE invoiceops_deliveries_total gauge",
        ]
    )
    for state, count in session.execute(
        select(Delivery.state, func.count()).group_by(Delivery.state)
    ):
        lines.append(f'invoiceops_deliveries_total{{state="{state}"}} {count}')
    duplicate_count = session.scalar(
        select(func.count())
        .select_from(AuditEvent)
        .where(AuditEvent.action == "exact_duplicate_detected")
    ) or 0
    lines.extend(
        [
            "# HELP invoiceops_exact_duplicates_total Exact duplicate upload attempts.",
            "# TYPE invoiceops_exact_duplicates_total counter",
            f"invoiceops_exact_duplicates_total {duplicate_count}",
            "# HELP invoiceops_processing_duration_milliseconds_sum Completed attempt time.",
            "# TYPE invoiceops_processing_duration_milliseconds_sum counter",
        ]
    )
    duration_sum = session.scalar(
        select(func.coalesce(func.sum(ProcessingAttempt.duration_ms), 0))
    ) or 0
    lines.append(f"invoiceops_processing_duration_milliseconds_sum {duration_sum}")
    return "\n".join(lines) + "\n"


def _decimal_export(value: object) -> Decimal | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError("Stored normalized amounts must be decimal strings.")
    try:
        return Decimal(value)
    except InvalidOperation as exc:
        raise ValueError("Stored normalized amount is invalid.") from exc


def _export_columns() -> list[str]:
    return [
        "document_id",
        "revision_id",
        "document_name",
        "sha256",
        "document_type",
        "supplier",
        "tax_identifier",
        "document_number",
        "issue_date",
        "currency",
        "subtotal",
        "discounts",
        "taxes",
        "shipping",
        "total",
        "approved_at",
    ]


def _document_json(document: Document) -> dict[str, object]:
    return {
        "id": str(document.id),
        "original_name": document.original_name,
        "content_type": document.content_type,
        "sha256": document.sha256,
        "size_bytes": document.size_bytes,
        "processing_status": document.processing_status,
        "review_status": document.review_status,
        "delivery_status": document.delivery_status,
        "version": document.version,
        "last_error": document.last_error,
        "created_at": document.created_at.isoformat(),
        "updated_at": document.updated_at.isoformat(),
    }


def _revision_json(revision: Revision) -> dict[str, object]:
    return {
        "id": str(revision.id),
        "number": revision.number,
        "kind": revision.kind,
        "source_pages": revision.source_pages,
        "extracted": revision.extracted,
        "normalized": revision.normalized,
        "issues": revision.issues,
        "mode": revision.mode,
        "model": revision.model_name,
        "prompt_version": revision.prompt_version,
        "approved_at": revision.approved_at.isoformat() if revision.approved_at else None,
    }


def _delivery_json(delivery: Delivery) -> dict[str, object]:
    return {
        "id": str(delivery.id),
        "revision_id": str(delivery.revision_id),
        "payload_version": delivery.payload_version,
        "idempotency_key": delivery.idempotency_key,
        "state": delivery.state,
        "attempts": delivery.attempts,
        "response_status": delivery.response_status,
        "last_error": delivery.last_error,
        "created_at": delivery.created_at.isoformat(),
        "updated_at": delivery.updated_at.isoformat(),
    }


async def _outbox_loop() -> None:
    while True:
        try:
            await asyncio.to_thread(recover_stale_processing, session_factory)
            from invoiceops.worker import dispatch_pending_once, recover_stale_deliveries

            await asyncio.to_thread(recover_stale_deliveries, session_factory)
            await asyncio.to_thread(dispatch_pending_once)
        except OperationalError:
            logger.warning("database unavailable; outbox reconciliation will retry")
        await asyncio.sleep(3)


app = _create_app()
